# app.py
import streamlit as st
import os, io, json
from dotenv import load_dotenv
import pdfplumber
from docx import Document
import openai

# Load .env
load_dotenv(dotenv_path=os.path.join(os.path.dirname(__file__), ".env"))

# Set API key for OpenAI
openai.api_key = os.getenv("OPENAI_API_KEY")

if not openai.api_key:
    st.error("OPENAI_API_KEY not found in .env file")
    st.stop()
st.set_page_config(page_title="JD -> Cover Letter + Video Script", layout="wide")

# ------------------ Helpers ------------------
def read_file_text(uploaded_file):
    if uploaded_file is None:
        return ""
    name = uploaded_file.name.lower()
    data = uploaded_file.read()
    try:
        if name.endswith(".pdf"):
            text = []
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                for p in pdf.pages:
                    txt = p.extract_text()
                    if txt:
                        text.append(txt)
            return "\n".join(text)
        elif name.endswith(".docx"):
            doc = Document(io.BytesIO(data))
            return "\n".join([p.text for p in doc.paragraphs])
        else:
            return data.decode("utf-8", errors="ignore")
    except Exception:
        return data.decode("utf-8", errors="ignore")

def call_chat_model(messages, max_tokens=800, temperature=0, model="gpt-3.5-turbo"):
    resp = openai.ChatCompletion.create(
        model=model,
        messages=messages,
        max_tokens=max_tokens,
        temperature=temperature
    )
    return resp.choices[0].message["content"]

#-------------------------------------------------------

def extract_resume_skills(resume_text):
    system = {"role": "system", "content": "You are an expert resume parser."}
    user = {
        "role": "user",
        "content": f"""
        Extract the technical skills from this resume text.
        Output a JSON object with a single key "skills" containing a list of strings.
        
        Resume Text:
        {resume_text}
        
        Output only valid JSON.
        """
    }
    out = call_chat_model([system, user], max_tokens=300, temperature=0)
    try:
        return json.loads(out).get("skills", [])
    except:
        return []
    



# ------------------ Core Functions ------------------
def extract_requirements(jd_text):
    system = {"role":"system","content":"You are an expert technical recruiter and extractor. Return only valid JSON (no markdown)."}
    user = {
        "role":"user",
        "content":(
            "Job description below. Extract a JSON object with keys:\n"
            "title (string), responsibilities (array of short strings), must_have (array), nice_to_have (array), "
            "experience_years (string or null), skills (array), keywords (array).\n\n"
            f"Job description:\n{jd_text}\n\n"
            "IMPORTANT: Output only valid JSON."
        )
    }
    out = call_chat_model([system, user], max_tokens=800, temperature=0)
    try:
        return json.loads(out)
    except Exception:
        start = out.find("{")
        end = out.rfind("}")
        if start != -1 and end != -1:
            try:
                return json.loads(out[start:end+1])
            except Exception:
                return {"raw": out}
        return {"raw": out}
    
def normalize_requirements(requirements_json):
    system = {
        "role": "system",
        "content": (
            "You are an expert tech recruiter and software architect. "
            "Your job is to translate vague, metaphorical, or HR-fluff requirements "
            "into concrete, industry-standard skills, responsibilities, tools, and role titles. "
            "Avoid metaphors. Be specific and realistic."
        )
    }

    user = {
        "role": "user",
        "content": f"""
Given the following extracted job requirements JSON, rewrite it into a NORMALIZED, CONCRETE version.

Rules:
- Replace metaphors with real technical skills and responsibilities
- Infer the most likely real-world job title
- Fill in skills and keywords with real technologies and concepts
- Keep the same JSON structure
- Output ONLY valid JSON

Input JSON:
{json.dumps(requirements_json, ensure_ascii=False, indent=2)}
"""
    }

    out = call_chat_model([system, user], max_tokens=800, temperature=0)

    try:
        return json.loads(out)
    except Exception:
        start = out.find("{")
        end = out.rfind("}")
        if start != -1 and end != -1:
            try:
                return json.loads(out[start:end+1])
            except Exception:
                return {"raw": out}
        return {"raw": out}


def generate_cover_letter(profile_text, requirements_json, job_title, tone="professional"):
    system = {
        "role": "system",
        "content": (
            "You are an expert career coach and technical writer. "
            "You write high-impact cover letters that connect a candidate's SPECIFIC projects "
            "to the job requirements. "
            "DO NOT lie about seniority. If the candidate is a junior/intern, "
            "focus heavily on their projects, quick learning, and hands-on work. "
            "DO NOT use generic phrases like 'I am writing to apply'. "
            "Instead, start with a strong hook about the candidate's relevant work."
        )
    }

    user = {
        "role": "user",
        "content": f"""
Write a high-impact, evidence-based cover letter for the job below.

JOB TITLE: {job_title}

JOB REQUIREMENTS (Normalized):
{json.dumps(requirements_json, ensure_ascii=False, indent=2)}

CANDIDATE PROFILE:
{profile_text}

STRICT WRITING RULES:
1. **The Hook:** Start by mentioning a SPECIFIC project or achievement from the profile that matches the job's key skills (e.g., "I recently built an AI Agent using Python..."). Do not start with "I am writing to apply...".
2. **Evidence over Claims:** Don't just say "I have experience with APIs." Say "I built [Project Name] using FastAPI to handle [specific task]."
3. **Honesty:** If the JD asks for 5 years and the candidate has less, do not fake "years of experience." Frame it as "Intensive project experience" or "Proven ability to ship production-ready code."
4. **Tone:** {tone}. Confident, specific, and humble.
5. **Call to Action:** End with a specific request (e.g., "I’d love to walk you through my code for [Project Name].")

Output only the body of the letter (no headers like 'Subject:' or placeholders).
"""
    }

    out = call_chat_model([system, user], max_tokens=700, temperature=0.3)
    return out

def analyze_fit(profile_text, requirements_json):
    system = {
        "role": "system", 
        "content": (
            "You are a strict technical recruiter. Your job is to analyze the gap between "
            "a candidate's resume and a job description. "
            "You provide brutal but helpful feedback."
        )
    }
    
    user = {
        "role": "user",
        "content": f"""
        Compare this candidate to the job requirements and provide a gap analysis.
        
        Job Requirements: {json.dumps(requirements_json)}
        Candidate Profile: {profile_text}
        
        Output a JSON object with exactly these keys:
        1. "match_score": (integer 0-100)
        2. "missing_keywords": (array of strings, listing critical skills the candidate lacks)
        3. "project_idea": (string, a specific project idea they should build THIS WEEK to bridge the gap)
        4. "advice": (string, 1-2 sentences on how to position themselves despite the gaps)
        
        Output ONLY valid JSON.
        """
    }
    
    out = call_chat_model([system, user], max_tokens=500, temperature=0)
    
    try:
        return json.loads(out)
    except Exception:
        # Fallback manual parsing if the LLM adds markdown
        start = out.find("{")
        end = out.rfind("}")
        if start != -1 and end != -1:
            try:
                return json.loads(out[start:end+1])
            except Exception:
                return {"match_score": 0, "advice": "Error parsing AI response.", "project_idea": "N/A", "missing_keywords": []}
        return {"match_score": 0, "advice": "Error parsing AI response.", "project_idea": "N/A", "missing_keywords": []}


def generate_video_script(profile_text, requirements_json, job_title, name="Candidate"):
    system = {
        "role": "system",
        "content": "You are a scriptwriter for short professional self-introduction videos for tech candidates."
    }

    user = {
        "role": "user",
        "content": f"""
Create a 60-second professional self-introduction video script for a job application.

Rules:
- Use concrete, professional language (no metaphors)
- Reference relevant skills and experience from the profile and job requirements
- Include timestamps (e.g., 0:00–0:10, 0:10–0:30, etc.)
- Include short camera/shot suggestions in brackets
- End with a short, confident call-to-action line
- Keep it realistic and suitable for LinkedIn or company career pages

Job title: {job_title}

Normalized job requirements:
{json.dumps(requirements_json, ensure_ascii=False, indent=2)}

Candidate name: {name}

Candidate profile / resume:
{profile_text}

Output only the script text.
"""
    }

    out = call_chat_model([system, user], max_tokens=500, temperature=0.3)
    return out


# ------------------ Streamlit UI ------------------
st.title("AI Agent — JD → Tailored Cover Letter + Video Script")

with st.sidebar:
    st.markdown("## Upload inputs")
    jd_file = st.file_uploader("Upload Job Description (PDF / DOCX / TXT)", type=["pdf","docx","txt"])
    resume_file = st.file_uploader("Upload Your Resume (optional)", type=["pdf","docx","txt"])
    profile_text_input = st.text_area("Or paste your resume/profile text (optional)", height=200)
    name = st.text_input("Your full name (for video script)", "Your Name")
    job_title_override = st.text_input("Override job title (optional)", "")
    tone = st.selectbox("Cover letter tone", ["professional","friendly","direct","enthusiastic"])

st.markdown("### Steps")
if st.button("Generate tailored cover letter + video script"):
    with st.spinner("Parsing files and calling the model..."):
        jd_text = read_file_text(jd_file) if jd_file else ""
        resume_text = read_file_text(resume_file) if resume_file else ""
        profile_text = profile_text_input.strip() or resume_text
        
        if not jd_text:
            st.error("Please upload a job description to continue.")
        else:
            # 1. Extract Requirements
            requirements = extract_requirements(jd_text)
            st.subheader("Extracted requirements (raw)")
            with st.expander("Show Raw JSON"):
                st.json(requirements)

            # 2. Normalize Requirements
            normalized = normalize_requirements(requirements)
            
            # --- NEW: Extract Resume Skills for Comparison ---
            with st.spinner("Analyzing resume skills..."):
                resume_skills = extract_resume_skills(profile_text)
            
            # --- NEW: Skills Snapshot UI ---
            st.markdown("---")
            st.subheader("⚔️ Skills Snapshot: You vs. The Job")
            
            # Create two columns for side-by-side comparison
            s_col1, s_col2 = st.columns(2)
            
            with s_col1:
                st.markdown("#### 🧑‍💻 Your Resume Skills")
                # Display as clear tags
                if resume_skills:
                    st.write(", ".join([f"`{s}`" for s in resume_skills]))
                else:
                    st.warning("No specific skills detected in resume.")

            with s_col2:
                st.markdown("#### 🏢 Job Requirements")
                # Get skills from the normalized JD JSON
                jd_skills = normalized.get("skills", []) or normalized.get("keywords", [])
                if jd_skills:
                    st.write(", ".join([f"`{s}`" for s in jd_skills[:15]])) # Show top 15 to avoid clutter
                else:
                    st.warning("No specific skills detected in JD.")
            
            # -----------------------------------------------

            st.subheader("Normalized requirements (concrete & real-world)")
            with st.expander("Show Normalized JSON"):
                st.json(normalized)

            # 3. Analyze Fit (NEW SECTION)
            st.markdown("---")
            st.subheader("📊 Match Analysis & Strategy")
            
            with st.spinner("Analyzing your resume against the JD..."):
                analysis = analyze_fit(profile_text, normalized)

            # Display Score with dynamic color
            score = analysis.get("match_score", 0)
            if score >= 80:
                st.markdown(f"### Match Score: :green[{score}%] 🚀")
            elif score >= 50:
                st.markdown(f"### Match Score: :orange[{score}%] ⚠️")
            else:
                st.markdown(f"### Match Score: :red[{score}%] 🛑")

            # Display Analysis Columns
            col1, col2 = st.columns(2)
            with col1:
                st.info(f"**💡 Strategic Advice:**\n\n{analysis.get('advice')}")
            with col2:
                missing = analysis.get('missing_keywords', [])
                if missing:
                    st.error(f"**❌ Missing Keywords:**\n\n{', '.join(missing)}")
                else:
                    st.success("**✅ No critical keywords missing!**")
            
            # Project Recommendation
            st.success(f"**🛠 Recommended Project to Bridge the Gap:**\n\n*{analysis.get('project_idea')}*")
            st.markdown("---")

            # 4. Generate Cover Letter
            job_title = job_title_override.strip() or (normalized.get("title") if isinstance(normalized, dict) else "")
            
            with st.spinner("Drafting cover letter..."):
                cover = generate_cover_letter(profile_text, normalized, job_title, tone=tone)
            
            st.subheader("Tailored Cover Letter")
            st.text_area("Cover letter", cover, height=300)
            st.download_button("Download cover letter (.txt)", cover, file_name="cover_letter.txt")

            # 5. Generate Video Script
            with st.spinner("Writing video script..."):
                script = generate_video_script(profile_text, normalized, job_title, name=name)
            
            st.subheader("1-minute Video Script")
            st.text_area("Video script", script, height=260)
            st.download_button("Download video script (.txt)", script, file_name="video_script.txt")

st.markdown("---")
st.caption("Built with an LLM. Check outputs for factual accuracy and customize the cover letter before sending.")
